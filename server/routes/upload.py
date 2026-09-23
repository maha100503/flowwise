"""
File upload endpoint — accepts CSV, Excel, TXT, PDF and returns parsed JSON data.
"""

import csv
import io
import json
import os
import uuid
from pathlib import Path
from fastapi import APIRouter, UploadFile, File, HTTPException

router = APIRouter(tags=["upload"])

# Directory for persisting uploaded file data
UPLOAD_DIR = Path(__file__).resolve().parent.parent / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

# In-memory cache for uploaded file data (keyed by file_id)
uploaded_files: dict[str, dict] = {}


def _save_to_disk(file_id: str, data: dict) -> None:
    """Persist parsed file data to disk."""
    path = UPLOAD_DIR / f"{file_id}.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, default=str)


def _load_from_disk(file_id: str) -> dict | None:
    """Load parsed file data from disk if it exists."""
    path = UPLOAD_DIR / f"{file_id}.json"
    if not path.is_file():
        return None
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def parse_csv(content: bytes, filename: str) -> dict:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    rows = [row for row in reader]
    columns = list(rows[0].keys()) if rows else []
    return {"rows": rows, "columns": columns, "row_count": len(rows), "filename": filename, "format": "csv"}


def parse_excel(content: bytes, filename: str) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)

    # Find the header row — skip blank/summary rows at the top
    headers = None
    skipped = []
    for raw_row in rows_iter:
        non_empty = [c for c in raw_row if c is not None and str(c).strip() != ""]
        # A header row should have mostly non-empty string cells
        if len(non_empty) >= 2 and all(isinstance(c, str) for c in non_empty):
            headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(raw_row)]
            break
        skipped.append(raw_row)

    if headers is None:
        # Fallback: use first row as headers
        if skipped:
            first = skipped[0]
            headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(first)]
            skipped = skipped[1:]
        else:
            wb.close()
            return {"rows": [], "columns": [], "row_count": 0, "filename": filename, "format": "excel"}

    rows = []
    for row in rows_iter:
        rows.append({headers[i]: (cell if cell is not None else "") for i, cell in enumerate(row) if i < len(headers)})
    wb.close()
    return {"rows": rows, "columns": headers, "row_count": len(rows), "filename": filename, "format": "excel"}


def parse_txt(content: bytes, filename: str) -> dict:
    text = content.decode("utf-8-sig")
    lines = text.strip().splitlines()
    # Try CSV-style detection
    if lines and ("," in lines[0] or "\t" in lines[0]):
        sep = "\t" if "\t" in lines[0] else ","
        reader = csv.DictReader(io.StringIO(text), delimiter=sep)
        rows = [row for row in reader]
        columns = list(rows[0].keys()) if rows else []
        return {"rows": rows, "columns": columns, "row_count": len(rows), "filename": filename, "format": "csv"}
    return {"text": text, "lines": lines, "line_count": len(lines), "filename": filename, "format": "text"}


def parse_pdf(content: bytes, filename: str) -> dict:
    from PyPDF2 import PdfReader
    reader = PdfReader(io.BytesIO(content))
    pages = []
    full_text = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        pages.append({"page": i + 1, "text": text})
        full_text.append(text)
    return {"text": "\n\n".join(full_text), "pages": pages, "page_count": len(pages), "filename": filename, "format": "pdf"}


def parse_json_file(content: bytes, filename: str) -> dict:
    data = json.loads(content.decode("utf-8-sig"))
    if isinstance(data, list) and data and isinstance(data[0], dict):
        columns = list(data[0].keys())
        return {"rows": data, "columns": columns, "row_count": len(data), "filename": filename, "format": "json"}
    return {"data": data, "filename": filename, "format": "json"}


PARSERS = {
    ".csv": parse_csv,
    ".xlsx": parse_excel,
    ".xls": parse_excel,
    ".txt": parse_txt,
    ".tsv": parse_txt,
    ".pdf": parse_pdf,
    ".json": parse_json_file,
}

ALLOWED_EXTENSIONS = set(PARSERS.keys())
MAX_FILE_SIZE = 50 * 1024 * 1024  # 50 MB


@router.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    filename = file.filename or "unknown"
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(400, f"Unsupported file type: {ext}. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}")

    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(400, f"File too large. Max size: {MAX_FILE_SIZE // (1024*1024)} MB")

    try:
        parsed = PARSERS[ext](content, filename)
    except Exception as e:
        raise HTTPException(400, f"Failed to parse {filename}: {str(e)}")

    file_id = str(uuid.uuid4())[:8]
    uploaded_files[file_id] = parsed
    _save_to_disk(file_id, parsed)

    # Build summary for frontend (don't send all rows back in the upload response)
    summary = {
        "file_id": file_id,
        "filename": filename,
        "format": parsed.get("format"),
        "row_count": parsed.get("row_count"),
        "columns": parsed.get("columns"),
        "line_count": parsed.get("line_count"),
        "page_count": parsed.get("page_count"),
    }
    # Remove None values
    summary = {k: v for k, v in summary.items() if v is not None}
    return summary


@router.get("/api/upload/{file_id}")
async def get_uploaded_file(file_id: str):
    if file_id not in uploaded_files:
        # Try loading from disk
        data = _load_from_disk(file_id)
        if data is None:
            raise HTTPException(404, "File not found or expired")
        uploaded_files[file_id] = data
    return uploaded_files[file_id]
